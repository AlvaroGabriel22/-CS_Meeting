"""The fifteen acceptance tests of Sprint 0, in the order the spec lists them.

Each one pins a promise the weekly workflow depends on.  They run against the
provisional fixtures described in ``tests/fixtures/build_fixtures.py``.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl.utils import get_column_letter, range_boundaries

from app.excel import parse_file
from app.excel.model import CellRole, SemanticType, TableShape, ValueType


def _iqc(fixture_files: dict[str, Path], dataset: str = "a"):
    return parse_file(fixture_files[f"iqc_dataset_{dataset}.xlsx"], "IQC").tables[0]


# --------------------------------------------------------------------------- #
# 1. Detect the table
# --------------------------------------------------------------------------- #
def test_01_table_is_detected_without_being_told_where_it_is(fixture_files) -> None:
    workbook = parse_file(fixture_files["iqc_dataset_a.xlsx"], "IQC")
    assert len(workbook.tables) == 1
    table = workbook.tables[0]
    assert table.shape is TableShape.MATRIX
    # the range is *discovered* and kept only as provenance
    assert table.source_range.startswith("B1:")
    assert table.header_row_count == 1 and table.label_col_count == 3


def test_01b_two_tables_in_one_sheet_are_separated(fixture_files) -> None:
    tables = parse_file(fixture_files["field_asr_casr.xlsx"], "FIELD").tables
    assert len(tables) == 2
    assert tables[0].source_range != tables[1].source_range
    assert [table.title for table in tables] == ["ASR — Field Quality", "CASR — Field Quality"]


# --------------------------------------------------------------------------- #
# 2. Detect merged cells
# --------------------------------------------------------------------------- #
def test_02_merged_cells_are_detected_and_kept(fixture_files) -> None:
    table = _iqc(fixture_files)
    assert table.merged_ranges, "the section/group merges must be reported"

    merged = [cell for cell in table.cells if cell.merged_range]
    anchors = [cell for cell in merged if cell.is_merge_anchor]
    assert merged and anchors
    # every merged cell names its range, and the range is a real rectangle
    for cell in merged[:20]:
        c1, r1, c2, r2 = range_boundaries(cell.merged_range)
        assert r2 >= r1 and c2 >= c1

    # the merged "SEC" label reaches every row it covers
    sections = {row.category for row in table.rows if row.category}
    assert sections == {"SEC", "TNP", "TECPLAM"}


# --------------------------------------------------------------------------- #
# 3-5. Years, months, weeks
# --------------------------------------------------------------------------- #
def test_03_years_are_detected(fixture_files) -> None:
    table = _iqc(fixture_files)
    years = [period for period in table.periods if period.kind.value == "year"]
    assert [period.year for period in years] == [2025, 2026]


def test_04_months_are_detected(fixture_files) -> None:
    table = _iqc(fixture_files)
    months = [period for period in table.periods if period.kind.value == "month"]
    assert [period.month for period in months] == [1, 2, 3, 4, 5, 6, 7, 8]
    assert [period.label for period in months][:3] == ["Jan", "Feb", "Mar"]


def test_05_weeks_are_detected(fixture_files) -> None:
    table = _iqc(fixture_files)
    weeks = [period for period in table.periods if period.kind.value == "week"]
    assert [period.week for period in weeks] == [31, 32]
    assert all(period.sort_key.endswith(f"W{period.week:02d}") for period in weeks)


# --------------------------------------------------------------------------- #
# 6-7. The window moves — same code, no edits
# --------------------------------------------------------------------------- #
def test_06_w32_becomes_w33(fixture_files) -> None:
    dataset_a = _iqc(fixture_files, "a")
    dataset_b = _iqc(fixture_files, "b")
    assert [p.week for p in dataset_a.periods if p.kind.value == "week"] == [31, 32]
    assert [p.week for p in dataset_b.periods if p.kind.value == "week"] == [33, 34]
    # structure read identically on both
    assert dataset_a.header_row_count == dataset_b.header_row_count
    assert dataset_a.label_col_count == dataset_b.label_col_count
    assert dataset_a.hierarchy == dataset_b.hierarchy


def test_07_w33_becomes_w34(fixture_files) -> None:
    dataset_b = _iqc(fixture_files, "b")
    dataset_c = _iqc(fixture_files, "c")
    weeks_b = {p.week for p in dataset_b.periods if p.kind.value == "week"}
    weeks_c = {p.week for p in dataset_c.periods if p.kind.value == "week"}
    assert 34 in weeks_b and 34 in weeks_c
    assert weeks_b == weeks_c == {33, 34}


# --------------------------------------------------------------------------- #
# 8-10. Values survive the trip
# --------------------------------------------------------------------------- #
def test_08_na_is_preserved(fixture_files) -> None:
    table = _iqc(fixture_files)
    nas = [cell for cell in table.cells if cell.value_type is ValueType.NA]
    assert nas, "the 2025 column of TECPLAM is NA"
    for cell in nas:
        assert cell.number is None  # never coerced to 0
        assert cell.raw_value == "NA"
        assert cell.role is CellRole.VALUE


def test_09_division_errors_are_preserved(fixture_files) -> None:
    table = _iqc(fixture_files)
    errors = [cell for cell in table.cells if cell.value_type is ValueType.ERROR]
    assert errors
    for cell in errors:
        assert cell.error_code == "#DIV/0!"
        assert cell.raw_value == "#DIV/0!"
        assert cell.number is None


def test_10_numbers_keep_original_and_gain_an_interpretation(fixture_files) -> None:
    table = _iqc(fixture_files)
    numbers = [
        cell
        for cell in table.cells
        if cell.role is CellRole.VALUE and cell.value_type is ValueType.NUMBER
    ]
    assert numbers
    for cell in numbers:
        assert cell.raw_value is not None  # original
        assert cell.number == float(cell.raw_value)  # interpreted
        assert cell.display_value  # rendered
    # thousands separators are a rendering, not a mutation
    large = [cell for cell in numbers if cell.number and cell.number >= 1000]
    assert large and any("," in cell.display_value for cell in large)


# --------------------------------------------------------------------------- #
# 11. Styles
# --------------------------------------------------------------------------- #
def test_11_styles_are_preserved_as_visual_metadata(fixture_files) -> None:
    table = _iqc(fixture_files)
    assert table.styles, "styles are deduplicated per table"

    header_cells = [cell for cell in table.cells if cell.role is CellRole.HEADER and cell.style_id]
    header_styles = [table.styles[cell.style_id] for cell in header_cells]
    assert any(style.bold for style in header_styles)
    assert any(style.fill_color for style in header_styles)
    assert any(style.borders for style in header_styles)

    widths = [column.width for column in table.columns if column.width]
    assert len(set(widths)) > 1, "different column widths must survive"
    heights = [row.height for row in table.rows if row.height]
    assert heights and len(set(heights)) > 1


# --------------------------------------------------------------------------- #
# 12. Original coordinates
# --------------------------------------------------------------------------- #
def test_12_source_coordinates_are_preserved(fixture_files) -> None:
    table = _iqc(fixture_files)
    c1, r1, _c2, _r2 = range_boundaries(table.source_range)
    for cell in table.cells:
        assert cell.source_address == f"{get_column_letter(c1 + cell.col)}{r1 + cell.row}"
    assert all(column.source_column for column in table.columns)
    assert all(row.source_row for row in table.rows)


# --------------------------------------------------------------------------- #
# 13. Hierarchy
# --------------------------------------------------------------------------- #
def test_13_hierarchy_category_subcategory_metric(fixture_files) -> None:
    table = _iqc(fixture_files)
    assert table.hierarchy == ("category", "subcategory", "metric")
    data_rows = [row for row in table.rows if not row.is_header_row]
    assert (data_rows[0].category, data_rows[0].subcategory, data_rows[0].metric) == (
        "SEC",
        "Total",
        "PPM",
    )
    assert {row.subcategory for row in data_rows} == {"Total", "TSI", "Packing"}
    assert {row.metric for row in data_rows} == {"PPM", "Def.", "Insp."}
    assert all(row.semantic is SemanticType.METRIC for row in data_rows)


def test_13b_field_hierarchy_indicator_model_series(fixture_files) -> None:
    """ASR → MX → Target/Result: the innermost level is a *series*, not a metric."""
    asr, casr = parse_file(fixture_files["field_asr_casr.xlsx"], "FIELD").tables
    for table, indicator in ((asr, "ASR"), (casr, "CASR")):
        assert table.hierarchy == ("category", "subcategory", "series")
        rows = [row for row in table.rows if not row.is_header_row]
        assert {row.category for row in rows} == {indicator}
        assert {row.subcategory for row in rows} == {"MX", "Mobile"}
        assert {row.series_type for row in rows} == {"Target", "Result"}
        assert {row.metric for row in rows} == {None}  # never a plain metric
        assert all(row.semantic is SemanticType.SERIES for row in rows)


# --------------------------------------------------------------------------- #
# 14-15. Variable width and variable period count
# --------------------------------------------------------------------------- #
def test_14_variable_number_of_columns(fixture_files) -> None:
    dataset_a = _iqc(fixture_files, "a")
    dataset_c = _iqc(fixture_files, "c")
    assert dataset_c.col_count == dataset_a.col_count + 1  # September
    assert dataset_c.row_count > dataset_a.row_count  # an extra metric per group
    assert dataset_c.label_col_count == dataset_a.label_col_count
    # the added row is a Target *series*, and it never pollutes the metrics
    assert "Target" in {row.series_type for row in dataset_c.rows if row.series_type}
    assert "Target" not in {row.metric for row in dataset_c.rows if row.metric}


def test_15_variable_number_of_periods(fixture_files) -> None:
    counts = {
        dataset: len(_iqc(fixture_files, dataset).periods) for dataset in ("a", "b", "c")
    }
    assert counts["a"] == counts["b"] == 12  # 2 years + 8 months + 2 weeks
    assert counts["c"] == 13  # September joined
    for dataset in ("a", "b", "c"):
        table = _iqc(fixture_files, dataset)
        assert all(column.period for column in table.columns if not column.is_label_column)
