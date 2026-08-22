"""FIELD: the real workbook, read as it is written.

The file arrived after IQC and it is a different shape — one sheet, one table,
a ``Target`` row and a ``Result`` row under each model — so these tests exist
to prove the system reads *that* file, and not a structure someone imagined for
it (ADR-0011).

They skip when the confidential workbook is not on the machine.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from app.excel import parse_file
from app.services import charts
from app.services.interpretation import from_normalized
from app.services.render_model import build_table_view


def _tables(path: Path):
    return [from_normalized(table, "FIELD") for table in parse_file(path, "FIELD").tables]


def _upload(client, path: Path):
    return client.post(
        "/api/uploads",
        data={"department": "FIELD", "createVersion": "true"},
        files={
            "file": (
                path.name,
                path.read_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    ).json()


# --------------------------------------------------------------------------- #
# The table — fidelity
# --------------------------------------------------------------------------- #
def test_the_structure_is_read_from_the_file(field_real: Path) -> None:
    tables = _tables(field_real)
    assert len(tables) == 1

    table = tables[0]
    assert table.source_range == "B3:R13"
    assert list(table.hierarchy) == ["category", "subcategory", "series"]
    assert [column.label for column in table.columns if column.period] == [
        "2025", "2026", "Jan", "Feb", "Mar", "Apr", "May", "Jun",
        "Jul", "Aug", "Sep", "Oct", "Nov", "Dec",
    ]

    models = [
        (row.category, row.subcategory, row.series_type)
        for row in table.rows
        if row.series_type
    ]
    assert models == [
        ("ASR", "MX", "Target"),
        ("ASR", "MX", "Result"),
        ("ASR", "Mobile", "Target"),
        ("ASR", "Mobile", "Result"),
        ("ASR", "APS", "Target"),
        ("ASR", "APS", "Result"),
        ("CASR", "Mobile", "Target"),
        ("CASR", "Mobile", "Result"),
    ]


def test_the_qualifier_row_annotates_the_periods_without_renaming_them(
    field_real: Path,
) -> None:
    """``Simulation / Result / Partial`` says how firm a figure is.

    It is not a second series axis — the Target/Result split lives in the rows
    — and it is not the period's name either.  Read as either one, the chart
    would plot a column called ``Simulation`` (ADR-0045).
    """
    table = _tables(field_real)[0]
    by_letter = {column.source_column: column for column in table.columns}

    assert by_letter["F"].label == "2026"
    assert by_letter["G"].label == "Jan"
    assert all(column.series_type is None for column in table.columns)
    # the word is not lost: it stays on the period as the token it came from
    assert "Simulation" in by_letter["F"].period.tokens
    assert "Partial" in by_letter["N"].period.tokens


def test_the_qualifier_row_is_still_drawn_in_the_table(field_real: Path) -> None:
    """The analyst wrote it, so the reader sees it."""
    view = build_table_view(_tables(field_real)[0])
    written = [
        cell["text"]
        for row in view["rows"]
        if row["kind"] == "header"
        for cell in row["cells"]
        if cell["text"]
    ]
    assert "Simulation" in written and "Partial" in written


def test_a_year_in_the_header_is_not_grouped_into_thousands(field_real: Path) -> None:
    """``2025``, not ``2,025``: a year is a name, not a quantity."""
    view = build_table_view(_tables(field_real)[0])
    header_text = [
        cell["text"] for row in view["rows"] if row["kind"] == "header" for cell in row["cells"]
    ]
    assert "2025" in header_text and "2026" in header_text
    assert not any("," in text and text.replace(",", "").isdigit() for text in header_text)


def test_every_value_is_the_cell_the_workbook_holds(field_real: Path) -> None:
    sheet = openpyxl.load_workbook(field_real, data_only=True)["MX Field KPI"]
    view = build_table_view(_tables(field_real)[0])

    checked = 0
    for row in view["rows"]:
        for cell in row["cells"]:
            if cell["kind"] != "value" or cell["value"] is None:
                continue
            assert cell["value"] == pytest.approx(float(sheet[cell["source"]].value))
            checked += 1
    assert checked > 50


def test_na_stays_the_word_the_analyst_wrote(field_real: Path) -> None:
    view = build_table_view(_tables(field_real)[0])
    texts = [cell["text"] for row in view["rows"] for cell in row["cells"]]
    assert texts.count("NA") == 15  # APS target all year, CASR until March


# --------------------------------------------------------------------------- #
# The charts
# --------------------------------------------------------------------------- #
def test_one_chart_per_model_result_against_target(field_real: Path) -> None:
    built = charts.build_charts(_tables(field_real), department="FIELD")

    assert built["metric"] is None, "FIELD names no metric, and none is invented"
    assert [(chart["category"], chart["subcategory"]) for chart in built["charts"]] == [
        ("ASR", "MX"),
        ("ASR", "Mobile"),
        ("ASR", "APS"),
        ("CASR", "Mobile"),
    ]
    for chart in built["charts"]:
        assert chart["kind"] == "pair"
        assert [series["label"] for series in chart["bars"]] == ["Result"]
        assert chart["line"]["label"] == "Target"
        assert chart["stacked"] is False


def test_the_first_model_of_each_category_is_the_one_shown(field_real: Path) -> None:
    """ASR·MX and CASR·Mobile by default; the rest wait in the configuration."""
    built = charts.build_charts(_tables(field_real), department="FIELD")
    shown = [chart["id"] for chart in built["charts"] if chart["enabled"]]
    assert shown == [
        "MX Field KPI · ASR · MX",
        "MX Field KPI · CASR · Mobile",
    ]


def test_the_plotted_numbers_are_the_workbook_s(field_real: Path) -> None:
    sheet = openpyxl.load_workbook(field_real, data_only=True)["MX Field KPI"]
    for chart in charts.build_charts(_tables(field_real), department="FIELD")["charts"]:
        for series in chart["bars"] + [chart["line"]]:
            for point in series["points"]:
                if point["value"] is None:
                    continue
                assert point["value"] == pytest.approx(float(sheet[point["source"]].value))


def test_na_is_a_gap_on_the_chart_never_a_zero(field_real: Path) -> None:
    """The APS target was never set; drawing it at zero would invent a target."""
    chart = next(
        chart
        for chart in charts.build_charts(_tables(field_real), department="FIELD")["charts"]
        if (chart["category"], chart["subcategory"]) == ("ASR", "APS")
    )
    months = [
        point["value"]
        for point, period in zip(chart["line"]["points"], chart["periods"])
        if period["kind"] == "month"
    ]
    assert months and all(value is None for value in months)


def test_the_line_is_cut_where_years_become_months(field_real: Path) -> None:
    """2026 and January are two blocks; a segment between them is not a trend."""
    for chart in charts.build_charts(_tables(field_real), department="FIELD")["charts"]:
        assert chart["breaks"] == [2]
        assert [period["kind"] for period in chart["periods"][:3]] == ["year", "year", "month"]


# --------------------------------------------------------------------------- #
# Through the API
# --------------------------------------------------------------------------- #
def test_the_charts_endpoint_serves_the_field_pairs(client, field_real: Path) -> None:
    created = _upload(client, field_real)
    body = client.get(f"/api/versions/{created['versionId']}/charts").json()

    assert body["department"] == "FIELD"
    assert body["metric"] is None
    assert len(body["charts"]) == 4
    assert sum(1 for chart in body["charts"] if chart["enabled"]) == 2


def test_the_presenter_can_switch_a_model_on(client, field_real: Path) -> None:
    version_id = _upload(client, field_real)["versionId"]
    client.put(
        "/api/departments/FIELD/settings",
        json={"chartSeries": {"MX Field KPI · ASR · Mobile": {"enabled": True}}},
    )

    body = client.get(f"/api/versions/{version_id}/charts").json()
    shown = [chart["id"] for chart in body["charts"] if chart["enabled"]]
    assert "MX Field KPI · ASR · Mobile" in shown
    assert len(shown) == 3


def test_a_model_can_be_switched_off(client, field_real: Path) -> None:
    version_id = _upload(client, field_real)["versionId"]
    client.put(
        "/api/departments/FIELD/settings",
        json={"chartSeries": {"MX Field KPI · ASR · MX": {"enabled": False}}},
    )

    body = client.get(f"/api/versions/{version_id}/charts").json()
    shown = [chart["id"] for chart in body["charts"] if chart["enabled"]]
    assert shown == ["MX Field KPI · CASR · Mobile"]


def test_choosing_rows_does_not_switch_a_chart_off(client, field_real: Path) -> None:
    """A composition that only picks rows says nothing about being shown."""
    version_id = _upload(client, field_real)["versionId"]
    body = client.get(f"/api/versions/{version_id}/charts").json()
    chart = next(item for item in body["charts"] if item["id"].endswith("ASR · MX"))
    target = next(
        option["key"] for option in chart["available"] if option["path"].endswith("Target")
    )

    result = next(
        option["key"] for option in chart["available"] if option["path"].endswith("Result")
    )
    client.put(
        "/api/departments/FIELD/settings",
        json={"chartSeries": {chart["id"]: {"bars": [target], "line": result}}},
    )

    again = client.get(f"/api/versions/{version_id}/charts").json()
    chosen = next(item for item in again["charts"] if item["id"] == chart["id"])
    assert chosen["enabled"] is True, "picking rows is not the same as hiding a chart"
    # the pair can even be read the other way round, if that is what the
    # presenter wants to show
    assert [series["label"] for series in chosen["bars"]] == ["Target"]
    assert chosen["line"]["label"] == "Result"


def test_a_chart_can_reach_across_models_and_says_which_is_which(
    client, field_real: Path
) -> None:
    """The presenter is not fenced into one model, and the legend keeps up."""
    version_id = _upload(client, field_real)["versionId"]
    body = client.get(f"/api/versions/{version_id}/charts").json()
    chart = next(item for item in body["charts"] if item["id"].endswith("ASR · MX"))
    options = {option["path"]: option["key"] for option in chart["available"]}

    client.put(
        "/api/departments/FIELD/settings",
        json={
            "chartSeries": {
                chart["id"]: {
                    "bars": [options["ASR · MX · Result"], options["ASR · Mobile · Result"]],
                    "line": options["ASR · MX · Target"],
                }
            }
        },
    )

    again = client.get(f"/api/versions/{version_id}/charts").json()
    chosen = next(item for item in again["charts"] if item["id"] == chart["id"])
    assert [series["label"] for series in chosen["bars"]] == ["MX · Result", "Mobile · Result"]
    assert chosen["line"]["label"] == "MX · Target"
